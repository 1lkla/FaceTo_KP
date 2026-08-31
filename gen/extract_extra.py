#!/usr/bin/env python3
"""把 h4m5t 库中的 pdf/docx/html/txt 提取成 markdown，存 data/source_extra/。

用 gen/.venv 运行：gen/.venv/bin/python3 gen/extract_extra.py
"""
import re
import shutil
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SRC_DIR = Path("/Users/1lk/workspace/40-knowledge/interview_prep/h4m5t_Sec-Interview")
OUT_DIR = PROJECT_ROOT / "data" / "source_extra"

SKIPPED = []   # (文件名, 原因)


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages):
        t = page.extract_text() or ""
        parts.append(t)
    text = clean_text("\n\n".join(parts))
    return f"# {path.stem}\n\n{text}"


def extract_docx(path: Path) -> str:
    import docx
    doc = docx.Document(str(path))
    parts = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = (para.style.name or "") if para.style else ""
        m = re.match(r"Heading (\d)", style)
        if m:
            parts.append(f"{'#' * min(int(m.group(1)) + 1, 6)} {t}")
        else:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            parts.append("| " + " | ".join(cells) + " |")
    return f"# {path.stem}\n\n" + clean_text("\n\n".join(parts))


def extract_html(path: Path) -> str:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    parts = []
    for el in soup.find_all(["h1", "h2", "h3", "h4", "p", "li"]):
        t = el.get_text(" ", strip=True)
        if not t:
            continue
        if el.name.startswith("h"):
            parts.append(f"{'#' * (int(el.name[1]) + 1)} {t}")
        elif el.name == "li":
            parts.append(f"- {t}")
        else:
            parts.append(t)
    return f"# {path.stem}\n\n" + clean_text("\n\n".join(parts))


def quality_check(name: str, text: str) -> bool:
    """太乱/太短则跳过。返回 True 表示保留。"""
    body = text
    if len(body) < 300:
        SKIPPED.append((name, f"提取后仅 {len(body)} 字，内容太少"))
        return False
    # 可读字符占比太低（乱码/二进制残留）则跳过
    printable = sum(1 for c in body if c.isprintable() or c in "\n\t")
    if printable / max(len(body), 1) < 0.9:
        SKIPPED.append((name, "不可读字符占比过高，疑似乱码"))
        return False
    return True


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    done, skipped_static = [], []

    # 1. OWASP Top 10 2021 PDF
    pdf = SRC_DIR / "OWASP-TOP10-2021中文版V1.0发布.pdf"
    text = extract_pdf(pdf)
    out = OUT_DIR / "owasp_top10_2021.md"
    if quality_check(pdf.name, text):
        out.write_text(text, encoding="utf-8")
        done.append((out.name, len(text)))

    # 2. 所有 docx
    for f in sorted(SRC_DIR.glob("*.docx")):
        try:
            text = extract_docx(f)
        except Exception as e:
            SKIPPED.append((f.name, f"docx 解析失败: {e}"))
            continue
        out = OUT_DIR / (f.stem + ".md")
        if quality_check(f.name, text):
            out.write_text(text, encoding="utf-8")
            done.append((out.name, len(text)))

    # 3. html
    for f in sorted(SRC_DIR.glob("*.html")):
        text = extract_html(f)
        out = OUT_DIR / (f.stem + ".md")
        if quality_check(f.name, text):
            out.write_text(text, encoding="utf-8")
            done.append((out.name, len(text)))

    # 4. doc（旧格式，用 macOS 自带 textutil 转换）
    import subprocess
    for f in sorted(SRC_DIR.glob("*.doc")):
        try:
            r = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(f)],
                               capture_output=True, text=True, timeout=60)
            text = clean_text(r.stdout)
        except Exception as e:
            SKIPPED.append((f.name, f"textutil 转换失败: {e}"))
            continue
        out = OUT_DIR / (f.stem + ".md")
        if quality_check(f.name, text):
            out.write_text(f"# {f.stem}\n\n{text}", encoding="utf-8")
            done.append((out.name, len(text)))

    # 5. txt 直接复制；内容只是外链则跳过
    for f in sorted(SRC_DIR.glob("*.txt")):
        text = clean_text(f.read_text(encoding="utf-8", errors="replace"))
        if re.fullmatch(r"https?://\S*", text):
            SKIPPED.append((f.name, f"内容只是一个外链（{text}），无实际材料"))
            continue
        out = OUT_DIR / (f.stem + ".md")
        out.write_text(f"# {f.stem}\n\n{text}", encoding="utf-8")
        done.append((out.name, len(text)))

    print(f"输出目录: {OUT_DIR}")
    print(f"成功 {len(done)} 个文件：")
    for name, size in done:
        print(f"  {name} ({size} 字)")
    if SKIPPED:
        print(f"跳过 {len(SKIPPED)} 个：")
        for name, why in SKIPPED:
            print(f"  {name}: {why}")


if __name__ == "__main__":
    main()
